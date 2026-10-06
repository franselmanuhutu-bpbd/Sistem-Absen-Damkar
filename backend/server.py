from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, UploadFile, File, Query
from fastapi.responses import StreamingResponse, JSONResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional
import uuid
import io
import json
import logging
import calendar as _cal
import random
from datetime import datetime, timezone, timedelta, date as _date

import bcrypt
import jwt

from seed_data import SEED_EMPLOYEES, SEED_TEAMS

# ---------------------------------------------------------------------------
# Setup
# ---------------------------------------------------------------------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

app = FastAPI(title="Sistem Informasi Absensi DAMKAR Mimika")
api = APIRouter(prefix="/api")

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALG = "HS256"

STATUSES = ["HDR", "OFF", "SKT", "TK", "IZN", "DL"]
STATUS_LABEL = {"HDR": "Hadir", "OFF": "Off", "SKT": "Sakit", "TK": "Tanpa Keterangan", "IZN": "Izin", "DL": "Dinas Luar"}
STATUS_HEX = {"HDR": "16A34A", "OFF": "64748B", "SKT": "2563EB", "TK": "DC2626", "IZN": "D97706", "DL": "9333EA"}

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger("damkar")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(user: dict) -> str:
    payload = {
        "sub": user["id"],
        "email": user["email"],
        "role": user["role"],
        "exp": datetime.now(timezone.utc) + timedelta(days=7),
        "type": "access",
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALG)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def clean(doc: dict) -> dict:
    if doc:
        doc.pop("_id", None)
    return doc


def last_day_of_month(year: int, month: int) -> int:
    return _cal.monthrange(year, month)[1]


def month_range(start: str, end: str) -> List[str]:
    """Return list of 'YYYY-MM' from start to end inclusive."""
    sy, sm = int(start[:4]), int(start[5:7])
    ey, em = int(end[:4]), int(end[5:7])
    out = []
    y, m = sy, sm
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            m = 1
            y += 1
    return out


async def get_current_user(request: Request) -> dict:
    token = None
    auth = request.headers.get("Authorization", "")
    if auth.startswith("Bearer "):
        token = auth[7:]
    if not token:
        token = request.cookies.get("access_token")
    if not token:
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALG])
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")
    user = await db.users.find_one({"id": payload["sub"]})
    if not user or user.get("status") == "INACTIVE":
        raise HTTPException(status_code=401, detail="User tidak ditemukan / nonaktif")
    user = clean(user)
    user.pop("password_hash", None)
    return user


def require_roles(*roles):
    async def dep(user: dict = Depends(get_current_user)):
        if user["role"] not in roles:
            raise HTTPException(status_code=403, detail="Anda tidak memiliki akses untuk tindakan ini")
        return user
    return dep


async def write_audit(user: dict, action: str, **extra):
    doc = {
        "id": str(uuid.uuid4()),
        "timestamp": now_iso(),
        "user_email": user.get("email"),
        "user_name": user.get("name"),
        "action": action,
    }
    doc.update(extra)
    await db.audit_logs.insert_one(doc)


async def resolve_teams_for_date(date_str: str) -> dict:
    """Return {employee_id: team_id} active on a given date (YYYY-MM-DD)."""
    cursor = db.team_assignments.find({"start_date": {"$lte": date_str}})
    result = {}
    async for a in cursor:
        end = a.get("end_date")
        if end and end < date_str:
            continue
        # keep the assignment with the latest start_date
        prev = result.get(a["employee_id"])
        if prev is None or a["start_date"] > prev[1]:
            result[a["employee_id"]] = (a["team_id"], a["start_date"])
    return {k: v[0] for k, v in result.items()}


# ---------------------------------------------------------------------------
# Models
# ---------------------------------------------------------------------------
class LoginIn(BaseModel):
    email: EmailStr
    password: str


class EmployeeIn(BaseModel):
    nama: str
    nip: str = ""
    pangkat: str = ""
    jabatan: str = ""


class UserIn(BaseModel):
    name: str
    email: EmailStr
    password: Optional[str] = None
    role: str = "operator"
    status: str = "ACTIVE"


class AssignmentIn(BaseModel):
    employee_id: str
    team_id: str
    start_date: str
    end_date: Optional[str] = None


class BatchAttendanceIn(BaseModel):
    date: str
    employee_ids: List[str]
    status: str


# ---------------------------------------------------------------------------
# Auth routes
# ---------------------------------------------------------------------------
@api.post("/auth/login")
async def login(body: LoginIn):
    email = body.email.lower().strip()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(body.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau password salah")
    if user.get("status") == "INACTIVE":
        raise HTTPException(status_code=403, detail="Akun dinonaktifkan")
    user = clean(user)
    token = create_token(user)
    user.pop("password_hash", None)
    return {"access_token": token, "user": user}


@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api.post("/auth/logout")
async def logout(user: dict = Depends(get_current_user)):
    return {"ok": True}


# ---------------------------------------------------------------------------
# User management (admin)
# ---------------------------------------------------------------------------
@api.get("/users")
async def list_users(user: dict = Depends(require_roles("admin"))):
    users = await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(1000)
    return users


@api.post("/users")
async def create_user(body: UserIn, user: dict = Depends(require_roles("admin"))):
    email = body.email.lower().strip()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    if not body.password:
        raise HTTPException(status_code=400, detail="Password wajib diisi")
    if body.role not in ("admin", "operator", "viewer"):
        raise HTTPException(status_code=400, detail="Role tidak valid")
    doc = {
        "id": str(uuid.uuid4()),
        "name": body.name,
        "email": email,
        "password_hash": hash_password(body.password),
        "role": body.role,
        "status": "ACTIVE",
        "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    await write_audit(user, "Membuat user baru", detail=f"{body.name} ({body.role})")
    return clean({**doc, "password_hash": None})


@api.put("/users/{uid}")
async def update_user(uid: str, body: UserIn, user: dict = Depends(require_roles("admin"))):
    existing = await db.users.find_one({"id": uid})
    if not existing:
        raise HTTPException(status_code=404, detail="User tidak ditemukan")
    update = {"name": body.name, "role": body.role, "status": body.status, "email": body.email.lower().strip()}
    if body.password:
        update["password_hash"] = hash_password(body.password)
    await db.users.update_one({"id": uid}, {"$set": update})
    await write_audit(user, "Mengubah user", detail=body.name)
    return {"ok": True}


@api.delete("/users/{uid}")
async def deactivate_user(uid: str, user: dict = Depends(require_roles("admin"))):
    if uid == user["id"]:
        raise HTTPException(status_code=400, detail="Tidak dapat menonaktifkan akun sendiri")
    await db.users.update_one({"id": uid}, {"$set": {"status": "INACTIVE"}})
    await write_audit(user, "Menonaktifkan user")
    return {"ok": True}


# ---------------------------------------------------------------------------
# Teams
# ---------------------------------------------------------------------------
@api.get("/teams")
async def list_teams(user: dict = Depends(get_current_user)):
    teams = await db.teams.find({}, {"_id": 0}).sort("order", 1).to_list(100)
    return teams


# ---------------------------------------------------------------------------
# Employees
# ---------------------------------------------------------------------------
@api.get("/employees")
async def list_employees(
    search: str = "",
    status: str = "",
    date: Optional[str] = None,
    user: dict = Depends(get_current_user),
):
    q = {}
    if status in ("ACTIVE", "INACTIVE"):
        q["status"] = status
    if search:
        q["$or"] = [
            {"nama": {"$regex": search, "$options": "i"}},
            {"nip": {"$regex": search, "$options": "i"}},
        ]
    employees = await db.employees.find(q, {"_id": 0}).sort("no", 1).to_list(5000)
    ref_date = date or _date.today().isoformat()
    team_map = await resolve_teams_for_date(ref_date)
    teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
    for e in employees:
        tid = team_map.get(e["id"])
        e["current_team_id"] = tid
        e["current_team_name"] = teams.get(tid, {}).get("name") if tid else None
    return employees


@api.post("/employees")
async def create_employee(body: EmployeeIn, user: dict = Depends(require_roles("admin", "operator"))):
    last = await db.employees.find_one({}, sort=[("no", -1)])
    no = (last["no"] + 1) if last else 1
    doc = {
        "id": str(uuid.uuid4()),
        "no": no,
        "nama": body.nama,
        "nip": body.nip,
        "pangkat": body.pangkat,
        "jabatan": body.jabatan,
        "status": "ACTIVE",
        "created_at": now_iso(),
    }
    await db.employees.insert_one(doc)
    await write_audit(user, "Menambah pegawai", employee_name=body.nama, detail=body.nip)
    return clean(doc)


@api.put("/employees/{eid}")
async def update_employee(eid: str, body: EmployeeIn, user: dict = Depends(require_roles("admin", "operator"))):
    existing = await db.employees.find_one({"id": eid})
    if not existing:
        raise HTTPException(status_code=404, detail="Pegawai tidak ditemukan")
    await db.employees.update_one({"id": eid}, {"$set": {
        "nama": body.nama, "nip": body.nip, "pangkat": body.pangkat, "jabatan": body.jabatan,
    }})
    await write_audit(user, "Mengubah data pegawai", employee_name=body.nama)
    return {"ok": True}


@api.post("/employees/{eid}/status")
async def toggle_employee_status(eid: str, user: dict = Depends(require_roles("admin", "operator"))):
    existing = await db.employees.find_one({"id": eid})
    if not existing:
        raise HTTPException(status_code=404, detail="Pegawai tidak ditemukan")
    new_status = "INACTIVE" if existing.get("status") == "ACTIVE" else "ACTIVE"
    await db.employees.update_one({"id": eid}, {"$set": {"status": new_status}})
    await write_audit(user, f"Mengubah status pegawai menjadi {new_status}", employee_name=existing["nama"])
    return {"ok": True, "status": new_status}


@api.get("/employees/{eid}/assignments")
async def employee_assignments(eid: str, user: dict = Depends(get_current_user)):
    assigns = await db.team_assignments.find({"employee_id": eid}, {"_id": 0}).sort("start_date", -1).to_list(500)
    teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
    for a in assigns:
        a["team_name"] = teams.get(a["team_id"], {}).get("name")
    return assigns


@api.post("/employees/import")
async def import_employees(file: UploadFile = File(...), user: dict = Depends(require_roles("admin"))):
    import pandas as pd
    content = await file.read()
    try:
        df = pd.read_excel(io.BytesIO(content), dtype=str)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Gagal membaca Excel: {e}")
    df.columns = [str(c).strip().lower() for c in df.columns]

    def pick(row, *keys):
        for k in keys:
            for col in df.columns:
                if k in col:
                    val = row.get(col)
                    if val is not None and str(val) != "nan":
                        return str(val).strip()
        return ""

    last = await db.employees.find_one({}, sort=[("no", -1)])
    no = (last["no"] + 1) if last else 1
    inserted = 0
    for _, row in df.iterrows():
        nama = pick(row, "nama", "name")
        if not nama:
            continue
        doc = {
            "id": str(uuid.uuid4()), "no": no, "nama": nama,
            "nip": pick(row, "nip"), "pangkat": pick(row, "pangkat", "golongan"),
            "jabatan": pick(row, "jabatan"), "status": "ACTIVE", "created_at": now_iso(),
        }
        await db.employees.insert_one(doc)
        no += 1
        inserted += 1
    await write_audit(user, "Import data pegawai dari Excel", detail=f"{inserted} pegawai")
    return {"inserted": inserted}


# ---------------------------------------------------------------------------
# Assignments / Rolling
# ---------------------------------------------------------------------------
@api.get("/assignments")
async def list_assignments(user: dict = Depends(get_current_user)):
    assigns = await db.team_assignments.find({}, {"_id": 0}).sort("start_date", -1).to_list(5000)
    return assigns


@api.post("/assignments")
async def create_assignment(body: AssignmentIn, user: dict = Depends(require_roles("admin", "operator"))):
    """Assign/roll an employee to a team. Closes any open prior assignment the day before start."""
    emp = await db.employees.find_one({"id": body.employee_id})
    team = await db.teams.find_one({"id": body.team_id})
    if not emp or not team:
        raise HTTPException(status_code=404, detail="Pegawai / Regu tidak ditemukan")
    # close open assignments that start before the new start
    open_ones = db.team_assignments.find({"employee_id": body.employee_id, "end_date": None})
    prev_day = (_date.fromisoformat(body.start_date) - timedelta(days=1)).isoformat()
    async for a in open_ones:
        if a["start_date"] < body.start_date:
            await db.team_assignments.update_one({"id": a["id"]}, {"$set": {"end_date": prev_day}})
    doc = {
        "id": str(uuid.uuid4()),
        "employee_id": body.employee_id,
        "team_id": body.team_id,
        "start_date": body.start_date,
        "end_date": body.end_date,
        "created_at": now_iso(),
    }
    await db.team_assignments.insert_one(doc)
    await write_audit(user, "Penempatan / Rolling regu", employee_name=emp["nama"],
                      detail=f"{team['name']} mulai {body.start_date}")
    return clean(doc)


@api.delete("/assignments/{aid}")
async def delete_assignment(aid: str, user: dict = Depends(require_roles("admin"))):
    await db.team_assignments.delete_one({"id": aid})
    await write_audit(user, "Menghapus penempatan regu")
    return {"ok": True}


@api.get("/teams/{team_id}/members")
async def team_members(team_id: str, date: str = None, user: dict = Depends(get_current_user)):
    ref = date or _date.today().isoformat()
    team_map = await resolve_teams_for_date(ref)
    ids = [eid for eid, tid in team_map.items() if tid == team_id]
    emps = await db.employees.find({"id": {"$in": ids}, "status": "ACTIVE"}, {"_id": 0}).sort("no", 1).to_list(5000)
    return emps


# ---------------------------------------------------------------------------
# Attendance
# ---------------------------------------------------------------------------
@api.get("/attendance/roster")
async def attendance_roster(date: str, team_id: str, user: dict = Depends(get_current_user)):
    team_map = await resolve_teams_for_date(date)
    ids = [eid for eid, tid in team_map.items() if tid == team_id]
    emps = await db.employees.find({"id": {"$in": ids}, "status": "ACTIVE"}, {"_id": 0}).sort("no", 1).to_list(5000)
    records = await db.attendance.find({"date": date, "employee_id": {"$in": ids}}, {"_id": 0}).to_list(5000)
    status_map = {r["employee_id"]: r["status"] for r in records}
    for e in emps:
        e["status"] = status_map.get(e["id"])
    return emps


@api.post("/attendance/batch")
async def batch_attendance(body: BatchAttendanceIn, user: dict = Depends(require_roles("admin", "operator"))):
    if body.status not in STATUSES:
        raise HTTPException(status_code=400, detail="Status tidak valid")
    team_map = await resolve_teams_for_date(body.date)
    count = 0
    for eid in body.employee_ids:
        existing = await db.attendance.find_one({"employee_id": eid, "date": body.date})
        old = existing["status"] if existing else None
        if old == body.status:
            continue
        doc = {
            "employee_id": eid,
            "date": body.date,
            "status": body.status,
            "team_id": team_map.get(eid),
            "updated_at": now_iso(),
            "updated_by": user["email"],
        }
        if existing:
            await db.attendance.update_one({"employee_id": eid, "date": body.date}, {"$set": doc})
        else:
            doc["id"] = str(uuid.uuid4())
            doc["created_at"] = now_iso()
            await db.attendance.insert_one(doc)
        emp = await db.employees.find_one({"id": eid}, {"_id": 0, "nama": 1})
        await write_audit(user, "Mengubah absensi", employee_name=emp["nama"] if emp else eid,
                          date=body.date, old_status=old, new_status=body.status)
        count += 1
    return {"updated": count}


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------
@api.get("/dashboard")
async def dashboard(date: str = None, user: dict = Depends(get_current_user)):
    ref = date or _date.today().isoformat()
    teams = await db.teams.find({}, {"_id": 0}).sort("order", 1).to_list(100)
    team_map = await resolve_teams_for_date(ref)
    active_emps = await db.employees.find({"status": "ACTIVE"}, {"_id": 0, "id": 1}).to_list(5000)
    total_employees = len(active_emps)
    records = await db.attendance.find({"date": ref}, {"_id": 0}).to_list(10000)
    rec_map = {r["employee_id"]: r["status"] for r in records}

    def empty():
        return {s: 0 for s in STATUSES}

    totals = empty()
    per_team = {t["id"]: {"team": t, "members": 0, **empty()} for t in teams}
    for e in active_emps:
        tid = team_map.get(e["id"])
        if tid in per_team:
            per_team[tid]["members"] += 1
        st = rec_map.get(e["id"])
        if st in STATUSES:
            totals[st] += 1
            if tid in per_team:
                per_team[tid][st] += 1
    return {
        "date": ref,
        "total_employees": total_employees,
        "totals": totals,
        "per_team": list(per_team.values()),
    }


# ---------------------------------------------------------------------------
# Recap computation
# ---------------------------------------------------------------------------
async def _compute_recap(start_month: str, end_month: str, team_id: Optional[str]):
    months = month_range(start_month, end_month)
    start_date = f"{start_month}-01"
    ey, em = int(end_month[:4]), int(end_month[5:7])
    end_date = f"{end_month}-{last_day_of_month(ey, em):02d}"
    # team resolution representative = last day of end month
    team_map = await resolve_teams_for_date(end_date)
    teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
    employees = await db.employees.find({}, {"_id": 0}).sort("no", 1).to_list(5000)
    records = await db.attendance.find(
        {"date": {"$gte": start_date, "$lte": end_date}}, {"_id": 0}
    ).to_list(200000)

    # per employee totals + breakdown per month
    totals_by_emp = {e["id"]: {s: 0 for s in STATUSES} for e in employees}
    breakdown = {e["id"]: {m: {s: 0 for s in STATUSES} for m in months} for e in employees}
    for r in records:
        eid = r["employee_id"]
        st = r["status"]
        mo = r["date"][:7]
        if eid in totals_by_emp and st in STATUSES:
            totals_by_emp[eid][st] += 1
            if mo in breakdown[eid]:
                breakdown[eid][mo][st] += 1

    rows = []
    grand = {s: 0 for s in STATUSES}
    for e in employees:
        tid = team_map.get(e["id"])
        if team_id and tid != team_id:
            continue
        t = totals_by_emp[e["id"]]
        total = sum(t.values())
        if total == 0 and e.get("status") == "INACTIVE":
            continue
        for s in STATUSES:
            grand[s] += t[s]
        rows.append({
            "employee_id": e["id"], "no": e["no"], "nip": e["nip"], "nama": e["nama"],
            "jabatan": e["jabatan"], "pangkat": e["pangkat"],
            "team_id": tid, "regu": teams.get(tid, {}).get("name") if tid else "-",
            **t, "total": total,
            "breakdown": {m: breakdown[e["id"]][m] for m in months},
        })
    return {
        "months": months, "rows": rows, "grand_total": grand,
        "total_pegawai": len(rows),
        "period_label": _period_label(start_month, end_month),
    }


def _month_label(m: str) -> str:
    names = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
             "Agustus", "September", "Oktober", "November", "Desember"]
    return f"{names[int(m[5:7]) - 1]} {m[:4]}"


def _period_label(start: str, end: str) -> str:
    if start == end:
        return _month_label(start)
    return f"{_month_label(start)} – {_month_label(end)}"


@api.get("/recap/monthly")
async def recap_monthly(month: str, team_id: str = None, user: dict = Depends(get_current_user)):
    return await _compute_recap(month, month, team_id or None)


@api.get("/recap/period")
async def recap_period(start: str, end: str, team_id: str = None, user: dict = Depends(get_current_user)):
    if start > end:
        raise HTTPException(status_code=400, detail="Bulan mulai tidak boleh setelah bulan selesai")
    return await _compute_recap(start, end, team_id or None)


# ---------------------------------------------------------------------------
# Calendar
# ---------------------------------------------------------------------------
@api.get("/calendar")
async def calendar_view(month: str, team_id: str = None, user: dict = Depends(get_current_user)):
    y, m = int(month[:4]), int(month[5:7])
    ld = last_day_of_month(y, m)
    start_date = f"{month}-01"
    end_date = f"{month}-{ld:02d}"
    q = {"date": {"$gte": start_date, "$lte": end_date}}
    if team_id:
        q["team_id"] = team_id
    records = await db.attendance.find(q, {"_id": 0}).to_list(200000)
    days = {}
    for d in range(1, ld + 1):
        days[f"{month}-{d:02d}"] = {s: 0 for s in STATUSES}
    for r in records:
        if r["date"] in days and r["status"] in STATUSES:
            days[r["date"]][r["status"]] += 1
    return {"month": month, "days": days}


@api.get("/attendance/day")
async def attendance_day(date: str, team_id: str = None, user: dict = Depends(get_current_user)):
    q = {"date": date}
    if team_id:
        q["team_id"] = team_id
    records = await db.attendance.find(q, {"_id": 0}).to_list(10000)
    emp_ids = [r["employee_id"] for r in records]
    emps = {e["id"]: e for e in await db.employees.find({"id": {"$in": emp_ids}}, {"_id": 0}).to_list(5000)}
    teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
    out = []
    for r in records:
        e = emps.get(r["employee_id"], {})
        out.append({
            "nama": e.get("nama"), "nip": e.get("nip"), "jabatan": e.get("jabatan"),
            "regu": teams.get(r.get("team_id"), {}).get("name"), "status": r["status"],
        })
    out.sort(key=lambda x: (x.get("regu") or "", x.get("nama") or ""))
    return out


# ---------------------------------------------------------------------------
# Audit log
# ---------------------------------------------------------------------------
@api.get("/audit")
async def audit_log(limit: int = 200, user: dict = Depends(get_current_user)):
    logs = await db.audit_logs.find({}, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    return logs


# ---------------------------------------------------------------------------
# Backup
# ---------------------------------------------------------------------------
@api.get("/backup")
async def backup(user: dict = Depends(require_roles("admin"))):
    dump = {}
    for col in ["users", "employees", "teams", "team_assignments", "attendance", "audit_logs"]:
        docs = await db[col].find({}, {"_id": 0}).to_list(500000)
        if col == "users":
            for d in docs:
                d.pop("password_hash", None)
        dump[col] = docs
    dump["_meta"] = {"generated_at": now_iso(), "db": os.environ["DB_NAME"]}
    data = json.dumps(dump, indent=2, default=str).encode("utf-8")
    fname = f"backup_damkar_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    await write_audit(user, "Backup database")
    return StreamingResponse(io.BytesIO(data), media_type="application/json",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


# ---------------------------------------------------------------------------
# Export Excel
# ---------------------------------------------------------------------------
KOP = {
    "l1": "PEMERINTAH KABUPATEN MIMIKA",
    "l2": "DINAS PEMADAM KEBAKARAN DAN PENYELAMATAN",
    "l3": "BIDANG PEMADAM KEBAKARAN",
    "addr": "Jl. Cenderawasih Km. 2, Timika, Papua Tengah",
}


@api.get("/export/excel")
async def export_excel(
    start: str, end: str, team_id: str = None,
    include_breakdown: bool = True, include_detail: bool = False,
    user: dict = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    data = await _compute_recap(start, end, team_id or None)
    wb = Workbook()
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, color="FFFFFF")
    title_font = Font(bold=True, size=14, color="0F172A")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    def autofit(ws, ncols, start_row=1):
        for c in range(1, ncols + 1):
            mx = 10
            for r in range(start_row, ws.max_row + 1):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    mx = max(mx, len(str(v)) + 2)
            ws.column_dimensions[get_column_letter(c)].width = min(mx, 45)

    # ---- Sheet 1: Ringkasan ----
    ws = wb.active
    ws.title = "Ringkasan"
    cols = ["No", "NIP", "Nama", "Regu", "Jabatan"] + STATUSES + ["Total"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(row=1, column=1, value=f"{KOP['l1']} - {KOP['l2']}").font = title_font
    ws.cell(row=1, column=1).alignment = center
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws.cell(row=2, column=1, value="REKAP ABSENSI PEGAWAI PEMADAM KEBAKARAN").font = Font(bold=True, size=12)
    ws.cell(row=2, column=1).alignment = center
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols))
    ws.cell(row=3, column=1, value=f"Periode: {data['period_label']}   |   Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M')}")
    ws.cell(row=3, column=1).alignment = center

    hr = 5
    for i, c in enumerate(cols, 1):
        ws.cell(row=hr, column=i, value=c)
    style_header(ws, hr, len(cols))
    r = hr + 1
    for row in data["rows"]:
        vals = [row["no"], row["nip"], row["nama"], row["regu"], row["jabatan"]] + [row[s] for s in STATUSES] + [row["total"]]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border = border
            if 6 <= i <= 11:
                cell.alignment = center
                cell.fill = PatternFill("solid", fgColor=STATUS_HEX[STATUSES[i - 6]])
                cell.font = Font(color="FFFFFF", bold=True)
        r += 1
    # totals row
    ws.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for i, s in enumerate(STATUSES):
        cell = ws.cell(row=r, column=6 + i, value=data["grand_total"][s])
        cell.font = Font(bold=True)
        cell.alignment = center
        cell.border = border
    ws.cell(row=r, column=12, value=sum(data["grand_total"].values())).font = Font(bold=True)
    autofit(ws, len(cols), hr)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{hr}"

    # ---- Sheet 2: Breakdown Bulanan ----
    if include_breakdown and len(data["months"]) > 0:
        wb2 = wb.create_sheet("Breakdown Bulanan")
        bcols = ["No", "Nama", "Regu"]
        for m in data["months"]:
            for s in STATUSES:
                bcols.append(f"{_month_label(m).split()[0][:3]} {s}")
        bcols.append("Total")
        for i, c in enumerate(bcols, 1):
            wb2.cell(row=1, column=i, value=c)
        style_header(wb2, 1, len(bcols))
        rr = 2
        for row in data["rows"]:
            vals = [row["no"], row["nama"], row["regu"]]
            for m in data["months"]:
                for s in STATUSES:
                    vals.append(row["breakdown"][m][s])
            vals.append(row["total"])
            for i, v in enumerate(vals, 1):
                c = wb2.cell(row=rr, column=i, value=v)
                c.border = border
            rr += 1
        autofit(wb2, len(bcols))
        wb2.freeze_panes = "D2"

    # ---- Sheet 3: Detail Absensi ----
    if include_detail:
        wb3 = wb.create_sheet("Detail Absensi")
        sy, sm = int(start[:4]), int(start[5:7])
        ey, em = int(end[:4]), int(end[5:7])
        sd = f"{start}-01"
        ed = f"{end}-{last_day_of_month(ey, em):02d}"
        records = await db.attendance.find({"date": {"$gte": sd, "$lte": ed}}, {"_id": 0}).sort("date", 1).to_list(200000)
        emps = {e["id"]: e for e in await db.employees.find({}, {"_id": 0}).to_list(5000)}
        teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
        dcols = ["Tanggal", "NIP", "Nama", "Regu", "Status", "Keterangan"]
        for i, c in enumerate(dcols, 1):
            wb3.cell(row=1, column=i, value=c)
        style_header(wb3, 1, len(dcols))
        rr = 2
        for rec in records:
            e = emps.get(rec["employee_id"], {})
            if team_id and rec.get("team_id") != team_id:
                continue
            vals = [rec["date"], e.get("nip"), e.get("nama"),
                    teams.get(rec.get("team_id"), {}).get("name"),
                    rec["status"], STATUS_LABEL.get(rec["status"])]
            for i, v in enumerate(vals, 1):
                c = wb3.cell(row=rr, column=i, value=v)
                c.border = border
                if i == 5:
                    c.fill = PatternFill("solid", fgColor=STATUS_HEX.get(rec["status"], "FFFFFF"))
                    c.font = Font(color="FFFFFF", bold=True)
                    c.alignment = center
            rr += 1
        autofit(wb3, len(dcols))
        wb3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Rekap_Absensi_DAMKAR_{start}_{end}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


# ---------------------------------------------------------------------------
# Export PDF
# ---------------------------------------------------------------------------
@api.get("/export/pdf")
async def export_pdf(
    start: str, end: str, team_id: str = None,
    include_summary: bool = True, include_breakdown: bool = False, include_detail: bool = False,
    user: dict = Depends(get_current_user),
):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (BaseDocTemplate, PageTemplate, Frame, Table, TableStyle,
                                    Paragraph, Spacer)
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER

    data = await _compute_recap(start, end, team_id or None)
    buf = io.BytesIO()
    page = landscape(A4)
    doc = BaseDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                          topMargin=30 * mm, bottomMargin=18 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")

    gen_time = datetime.now().strftime("%d-%m-%Y %H:%M")

    def header_footer(canvas, d):
        canvas.saveState()
        w, h = page
        canvas.setFont("Helvetica-Bold", 13)
        canvas.drawCentredString(w / 2, h - 14 * mm, KOP["l1"])
        canvas.setFont("Helvetica-Bold", 11)
        canvas.drawCentredString(w / 2, h - 19 * mm, KOP["l2"])
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(w / 2, h - 23.5 * mm, KOP["l3"] + " — " + KOP["addr"])
        canvas.setLineWidth(1.2)
        canvas.line(12 * mm, h - 26 * mm, w - 12 * mm, h - 26 * mm)
        # footer
        canvas.setFont("Helvetica-Oblique", 7.5)
        canvas.drawString(12 * mm, 10 * mm, "Dicetak dari Sistem Informasi Absensi DAMKAR")
        canvas.drawCentredString(w / 2, 10 * mm, f"Dicetak: {gen_time}")
        canvas.drawRightString(w - 12 * mm, 10 * mm, f"Halaman {d.page}")
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=header_footer)])

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=13, alignment=TA_CENTER, spaceAfter=2)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=10, alignment=TA_CENTER, spaceAfter=2)
    sec_style = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=11, textColor=colors.HexColor("#0F172A"), spaceBefore=8, spaceAfter=4)

    elements = []
    elements.append(Paragraph("REKAP ABSENSI PEGAWAI PEMADAM KEBAKARAN", title_style))
    elements.append(Paragraph(f"Periode: {data['period_label']}", sub_style))
    elements.append(Spacer(1, 6))

    status_colors = {s: colors.HexColor("#" + STATUS_HEX[s]) for s in STATUSES}

    if include_summary:
        elements.append(Paragraph("RINGKASAN PERIODE", sec_style))
        gt = data["grand_total"]
        sum_data = [["Total Pegawai"] + [STATUS_LABEL[s] for s in STATUSES] + ["Total"],
                    [data["total_pegawai"]] + [gt[s] for s in STATUSES] + [sum(gt.values())]]
        st = Table(sum_data, repeatRows=1)
        st.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(st)
        elements.append(Spacer(1, 8))

        elements.append(Paragraph("REKAP PER PEGAWAI", sec_style))
        header = ["No", "NIP", "Nama", "Regu", "Jabatan"] + STATUSES + ["Total"]
        table_data = [header]
        for row in data["rows"]:
            table_data.append([row["no"], row["nip"], Paragraph(str(row["nama"]), styles["BodyText"]),
                               row["regu"], Paragraph(str(row["jabatan"]), styles["BodyText"])]
                              + [row[s] for s in STATUSES] + [row["total"]])
        total_row = ["", "", "TOTAL", "", ""] + [data["grand_total"][s] for s in STATUSES] + [sum(data["grand_total"].values())]
        table_data.append(total_row)
        colw = [10 * mm, 34 * mm, 55 * mm, 20 * mm, 55 * mm] + [13 * mm] * 6 + [15 * mm]
        t = Table(table_data, colWidths=colw, repeatRows=1)
        ts = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("ALIGN", (5, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
            ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("SPAN", (2, -1), (4, -1)),
        ]
        for i, s in enumerate(STATUSES):
            ts.append(("TEXTCOLOR", (5 + i, 1), (5 + i, -2), status_colors[s]))
            ts.append(("FONTNAME", (5 + i, 1), (5 + i, -2), "Helvetica-Bold"))
        t.setStyle(TableStyle(ts))
        elements.append(t)

    if include_breakdown and len(data["months"]) > 0:
        elements.append(Paragraph("BREAKDOWN BULANAN", sec_style))
        header = ["Nama", "Regu"]
        for m in data["months"]:
            for s in STATUSES:
                header.append(f"{_month_label(m).split()[0][:3]}\n{s}")
        bd = [header]
        for row in data["rows"]:
            vals = [Paragraph(str(row["nama"]), styles["BodyText"]), row["regu"]]
            for m in data["months"]:
                for s in STATUSES:
                    vals.append(row["breakdown"][m][s])
            bd.append(vals)
        nmonth = len(data["months"])
        colw = [45 * mm, 18 * mm] + [((page[0] - 24 * mm - 63 * mm) / (nmonth * 6))] * (nmonth * 6)
        t = Table(bd, colWidths=colw, repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 6.5),
            ("ALIGN", (1, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(t)

    if include_detail:
        elements.append(Paragraph("DETAIL ABSENSI", sec_style))
        ey, em = int(end[:4]), int(end[5:7])
        sd = f"{start}-01"
        ed = f"{end}-{last_day_of_month(ey, em):02d}"
        records = await db.attendance.find({"date": {"$gte": sd, "$lte": ed}}, {"_id": 0}).sort("date", 1).to_list(200000)
        emps = {e["id"]: e for e in await db.employees.find({}, {"_id": 0}).to_list(5000)}
        teams = {t["id"]: t for t in await db.teams.find({}, {"_id": 0}).to_list(100)}
        dh = [["Tanggal", "NIP", "Nama", "Regu", "Status"]]
        for rec in records:
            if team_id and rec.get("team_id") != team_id:
                continue
            e = emps.get(rec["employee_id"], {})
            dh.append([rec["date"], e.get("nip"), Paragraph(str(e.get("nama")), styles["BodyText"]),
                       teams.get(rec.get("team_id"), {}).get("name"), rec["status"]])
        t = Table(dh, colWidths=[24 * mm, 36 * mm, 70 * mm, 25 * mm, 20 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ALIGN", (4, 0), (4, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(t)

    doc.build(elements)
    buf.seek(0)
    fname = f"Rekap_Absensi_DAMKAR_{start}_{end}.pdf"
    return StreamingResponse(buf, media_type="application/pdf",
                             headers={"Content-Disposition": f"attachment; filename={fname}"})


# ---------------------------------------------------------------------------
# Startup: seed
# ---------------------------------------------------------------------------
async def seed():
    # indexes
    await db.users.create_index("email", unique=True)
    await db.employees.create_index("no")
    await db.team_assignments.create_index("employee_id")
    await db.attendance.create_index([("employee_id", 1), ("date", 1)], unique=True)
    await db.attendance.create_index("date")

    # admin
    admin_email = os.environ["ADMIN_EMAIL"].lower()
    admin_pw = os.environ["ADMIN_PASSWORD"]
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "id": str(uuid.uuid4()), "name": "Administrator DAMKAR", "email": admin_email,
            "password_hash": hash_password(admin_pw), "role": "admin", "status": "ACTIVE",
            "created_at": now_iso(),
        })
    elif not verify_password(admin_pw, existing["password_hash"]):
        await db.users.update_one({"email": admin_email}, {"$set": {"password_hash": hash_password(admin_pw)}})

    # demo operator + viewer
    for em, nm, role in [("operator@damkar.go.id", "Operator Piket", "operator"),
                         ("kepala@damkar.go.id", "Kepala Bidang", "viewer")]:
        if not await db.users.find_one({"email": em}):
            await db.users.insert_one({
                "id": str(uuid.uuid4()), "name": nm, "email": em,
                "password_hash": hash_password("Damkar2026!"), "role": role, "status": "ACTIVE",
                "created_at": now_iso(),
            })

    # teams
    team_ids = {}
    for t in SEED_TEAMS:
        existing = await db.teams.find_one({"code": t["code"]})
        if existing:
            team_ids[t["order"]] = existing["id"]
        else:
            tid = str(uuid.uuid4())
            await db.teams.insert_one({"id": tid, **t})
            team_ids[t["order"]] = tid

    # employees
    emp_ids = []
    if await db.employees.count_documents({}) == 0:
        for e in SEED_EMPLOYEES:
            eid = str(uuid.uuid4())
            emp_ids.append(eid)
            await db.employees.insert_one({
                "id": eid, "no": e["no"], "nama": e["nama"], "nip": e["nip"],
                "pangkat": e["pangkat"], "jabatan": e["jabatan"], "status": "ACTIVE",
                "created_at": now_iso(),
            })
    else:
        emp_ids = [e["id"] for e in await db.employees.find({}, {"_id": 0, "id": 1}).sort("no", 1).to_list(5000)]

    # initial assignments (distribute across 6 regu from 2026-01-01) + demo attendance
    if await db.team_assignments.count_documents({}) == 0 and emp_ids:
        order_list = sorted(team_ids.keys())
        for i, eid in enumerate(emp_ids):
            tid = team_ids[order_list[i % len(order_list)]]
            await db.team_assignments.insert_one({
                "id": str(uuid.uuid4()), "employee_id": eid, "team_id": tid,
                "start_date": "2026-01-01", "end_date": None, "created_at": now_iso(),
            })

    # demo attendance from Jan 2026 through current month (idempotent via unique index)
    if await db.attendance.count_documents({}) == 0 and emp_ids:
        rnd = random.Random(42)
        today = _date.today()
        last_month = today.month if today.year >= 2026 else 12
        team_map = await resolve_teams_for_date(today.isoformat())
        weights = (["HDR"] * 68 + ["OFF"] * 15 + ["SKT"] * 5 + ["IZN"] * 5 + ["DL"] * 5 + ["TK"] * 2)
        docs = []
        for month in range(1, last_month + 1):
            ld = last_day_of_month(2026, month)
            if month == today.month and today.year == 2026:
                ld = today.day
            for day in range(1, ld + 1):
                dstr = f"2026-{month:02d}-{day:02d}"
                for eid in emp_ids:
                    st = rnd.choice(weights)
                    docs.append({
                        "id": str(uuid.uuid4()), "employee_id": eid, "date": dstr,
                        "status": st, "team_id": team_map.get(eid),
                        "created_at": now_iso(), "updated_by": "seed",
                    })
        # insert in chunks
        for i in range(0, len(docs), 2000):
            try:
                await db.attendance.insert_many(docs[i:i + 2000], ordered=False)
            except Exception:
                pass

    logger.info("Seeding complete")


@app.on_event("startup")
async def on_startup():
    await seed()


@api.get("/")
async def root():
    return {"message": "DAMKAR Absensi API"}


app.include_router(api)
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get("CORS_ORIGINS", "*").split(","),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("shutdown")
async def shutdown():
    client.close()
