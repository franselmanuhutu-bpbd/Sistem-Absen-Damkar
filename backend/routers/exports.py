from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from datetime import datetime
from typing import Optional
import io

from config import STATUSES, STATUS_LABEL, STATUS_HEX
from database import get_db, fetch_all
from auth import get_current_user, require_roles
from utils import (
    last_day_of_month,
    _compute_recap,
    _compute_kasubid_recap,
)

router = APIRouter(tags=["Exports"])

KOP = {
    "l1": "BADAN PENANGGULANGAN BENCANA DAERAH",
    "l2": "KABUPATEN MIMIKA",
    "l3": "BIDANG PEMADAM KEBAKARAN",
    "addr": "Jl. Cenderawasih Km. 2, Timika, Papua Tengah",
}


def sanitize_excel_cell(v):
    """Prevent CSV / Formula Injection (CWE-1236) in generated spreadsheets."""
    if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")):
        return f"'{v}"
    return v


@router.get("/export/excel")
async def export_excel(
    start: str,
    end: str,
    team_id: str = None,
    category: str = None,
    include_breakdown: bool = True,
    include_detail: bool = False,
    user: dict = Depends(get_current_user),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    data = await _compute_recap(start, end, team_id or None, category or None)
    wb = Workbook()
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="0F172A")
    header_font = Font(bold=True, size=10, color="FFFFFF")
    title_font = Font(bold=True, size=13, color="0F172A")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)
    data_font = Font(size=10)
    data_font_bold = Font(size=10, bold=True)

    def style_header(ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = border

    def autofit(ws, ncols, start_row=1):
        for c in range(1, ncols + 1):
            mx = 10
            for r in range(start_row, ws.max_row + 1):
                v = ws.cell(row=r, column=c).value
                if v is not None:
                    mx = max(mx, len(str(v)) + 2)
            ws.column_dimensions[get_column_letter(c)].width = min(mx, 45)

    ws = wb.active
    ws.title = "Ringkasan"
    cols = ["No", "NIP", "Nama", "Regu", "Kategori", "Jabatan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cols))
    ws.cell(row=1, column=1, value=f"{KOP['l1']} {KOP['l2']} \u2014 {KOP['l3']}").font = title_font
    ws.cell(row=1, column=1).alignment = center
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(cols))
    ws.cell(row=2, column=1, value="REKAPITULASI ABSENSI PEGAWAI").font = Font(bold=True, size=11)
    ws.cell(row=2, column=1).alignment = center
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(cols))
    ws.cell(row=3, column=1, value=f"Periode: {data['period_label']}   |   Dicetak: {datetime.now().strftime('%d-%m-%Y %H:%M')}").font = Font(size=10)
    ws.cell(row=3, column=1).alignment = center

    hr = 5
    for i, c in enumerate(cols, 1):
        ws.cell(row=hr, column=i, value=c)
    style_header(ws, hr, len(cols))
    r = hr + 1
    for row in data["rows"]:
        vals = [row["no"], row["nip"], row["nama"], row["regu"], row["category"], row["jabatan"]] + [row[s] for s in STATUSES] + [row["jumlah_hari_kerja"], row["total_kehadiran"]]
        for i, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=i, value=sanitize_excel_cell(v))
            cell.border = border
            cell.font = data_font
            if i == 3:  # Nama: Left Center
                cell.alignment = left_center
            else:       # Other: Middle
                cell.alignment = center
            if 7 <= i <= 12:
                cell.fill = PatternFill("solid", fgColor=STATUS_HEX[STATUSES[i - 7]])
                cell.font = Font(size=10, color="FFFFFF", bold=True)
            elif i in (len(vals) - 1, len(vals)):
                cell.font = data_font_bold
        r += 1

    ws.cell(row=r, column=1, value="TOTAL").font = data_font_bold
    ws.cell(row=r, column=1).alignment = center
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
    for i, s in enumerate(STATUSES):
        cell = ws.cell(row=r, column=7 + i, value=data["grand_total"][s])
        cell.font = data_font_bold
        cell.alignment = center
        cell.border = border
    cell_jhk = ws.cell(row=r, column=len(cols) - 1, value=sum(row.get("jumlah_hari_kerja", 0) for row in data["rows"]))
    cell_jhk.font = data_font_bold
    cell_jhk.alignment = center
    cell_jhk.border = border
    cell_hdr = ws.cell(row=r, column=len(cols), value=data["grand_total"]["HDR"])
    cell_hdr.font = data_font_bold
    cell_hdr.alignment = center
    cell_hdr.border = border
    autofit(ws, len(cols), hr)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)
    ws.auto_filter.ref = f"A{hr}:{get_column_letter(len(cols))}{hr}"

    if include_breakdown and len(data["breakdown"]) > 0:
        wb2 = wb.create_sheet("Breakdown Bulanan")
        bcols = ["No", "Nama", "Regu", "Bulan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
        for i, c in enumerate(bcols, 1):
            wb2.cell(row=1, column=i, value=c)
        style_header(wb2, 1, len(bcols))
        rr = 2
        for b in data["breakdown"]:
            vals = [b["no"], b["nama"], b["regu"], b["month_label"]] + [b[s] for s in STATUSES] + [b["jumlah_hari_kerja"], b["total_kehadiran"]]
            for i, v in enumerate(vals, 1):
                c = wb2.cell(row=rr, column=i, value=sanitize_excel_cell(v))
                c.border = border
                c.font = data_font
                if i == 2:  # Nama: Left Center
                    c.alignment = left_center
                else:       # Other: Middle
                    c.alignment = center
                if 5 <= i <= 10:
                    c.fill = PatternFill("solid", fgColor=STATUS_HEX[STATUSES[i - 5]])
                    c.font = Font(size=10, color="FFFFFF", bold=True)
                elif 11 <= i <= 12:
                    c.font = data_font_bold
            rr += 1
        autofit(wb2, len(bcols))
        wb2.freeze_panes = "E2"

    if include_detail:
        wb3 = wb.create_sheet("Detail Absensi")
        db = await get_db()
        sy, sm = int(start[:4]), int(start[5:7])
        ey, em = int(end[:4]), int(end[5:7])
        sd = f"{start}-01"
        ed = f"{end}-{last_day_of_month(ey, em):02d}"
        q = db.table("attendance").select("*").gte("date", sd).lte("date", ed).order("date")
        records = await fetch_all(q)
        emps = {e["id"]: e for e in ((await db.table("employees").select("*").execute()).data or [])}
        teams = {t["id"]: t for t in ((await db.table("teams").select("*").execute()).data or [])}
        dcols = ["Tanggal", "NIP", "Nama", "Regu", "Status", "Keterangan"]
        for i, c in enumerate(dcols, 1):
            wb3.cell(row=1, column=i, value=c)
        style_header(wb3, 1, len(dcols))
        rr = 2
        for rec in records:
            e = emps.get(rec["employee_id"], {})
            if team_id and rec.get("team_id") != team_id:
                continue
            vals = [
                rec["date"], e.get("nip"), e.get("nama"),
                teams.get(rec.get("team_id"), {}).get("name"),
                rec["status"], STATUS_LABEL.get(rec["status"])
            ]
            for i, v in enumerate(vals, 1):
                c = wb3.cell(row=rr, column=i, value=sanitize_excel_cell(v))
                c.border = border
                c.font = data_font
                if i == 3:  # Nama: Left Center
                    c.alignment = left_center
                else:       # Other: Middle
                    c.alignment = center
                if i == 5:
                    c.fill = PatternFill("solid", fgColor=STATUS_HEX.get(rec["status"], "FFFFFF"))
                    c.font = Font(size=10, color="FFFFFF", bold=True)
            rr += 1
        autofit(wb3, len(dcols))
        wb3.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"Rekap_Absensi_DAMKAR_{start}_{end}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )


@router.get("/export/pdf")
async def export_pdf(
    start: str,
    end: str,
    team_id: str = None,
    category: str = None,
    include_summary: bool = True,
    include_breakdown: bool = False,
    include_detail: bool = False,
    user: dict = Depends(get_current_user),
):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        BaseDocTemplate, PageTemplate, Frame, Table, TableStyle,
        Paragraph, Spacer
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    data = await _compute_recap(start, end, team_id or None, category or None)
    buf = io.BytesIO()
    page = landscape(A4)
    doc = BaseDocTemplate(buf, pagesize=page, leftMargin=12 * mm, rightMargin=12 * mm,
                          topMargin=30 * mm, bottomMargin=18 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="main")
    gen_time = datetime.now().strftime("%d-%m-%Y %H:%M")

    def header_footer(canvas, d):
        canvas.saveState()
        w, h = page
        canvas.setFont("Helvetica-Bold", 13)
        canvas.drawCentredString(w / 2, h - 14 * mm, KOP["l1"])
        canvas.setFont("Helvetica-Bold", 11)
        canvas.drawCentredString(w / 2, h - 19 * mm, KOP["l2"])
        canvas.setFont("Helvetica", 9)
        canvas.drawCentredString(w / 2, h - 23.5 * mm, KOP["l3"] + " — " + KOP["addr"])
        canvas.setLineWidth(1.2)
        canvas.line(12 * mm, h - 26 * mm, w - 12 * mm, h - 26 * mm)
        canvas.setFont("Helvetica-Oblique", 7.5)
        canvas.drawString(12 * mm, 10 * mm, "Dicetak dari Sistem Informasi Absensi DAMKAR")
        canvas.drawCentredString(w / 2, 10 * mm, f"Dicetak: {gen_time}")
        canvas.drawRightString(w - 12 * mm, 10 * mm, f"Halaman {d.page}")
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=header_footer)])

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], fontSize=13, alignment=TA_CENTER, spaceAfter=2)
    sub_style = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, alignment=TA_CENTER, spaceAfter=2)
    sec_style = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=10, textColor=colors.HexColor("#0F172A"), spaceBefore=8, spaceAfter=4)
    
    CELL_FONT_SIZE = 7
    CELL_LEADING = 8.5

    table_header_style = ParagraphStyle(
        "table_header",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        textColor=colors.white,
        wordWrap="LTR",
    )

    table_cell_center_style = ParagraphStyle(
        "table_cell_center",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        wordWrap="LTR",
    )

    table_cell_center_bold = ParagraphStyle(
        "table_cell_center_bold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        wordWrap="LTR",
    )

    table_cell_nama_style = ParagraphStyle(
        "table_cell_nama",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_LEFT,
        wordWrap="LTR",
    )

    status_colors = {s: colors.HexColor("#" + STATUS_HEX[s]) for s in STATUSES}

    status_cell_styles = {
        s: ParagraphStyle(
            f"table_status_{s}",
            parent=styles["Normal"],
            textColor=status_colors[s],
            fontName="Helvetica-Bold",
            fontSize=CELL_FONT_SIZE,
            leading=CELL_LEADING,
            alignment=TA_CENTER,
            wordWrap="LTR",
        )
        for s in STATUSES
    }

    def format_regu(value):
        return str(value).replace(" - ", "<br/>")

    elements = []
    elements.append(Paragraph("REKAPITULASI ABSENSI PEGAWAI", title_style))
    elements.append(Paragraph(f"Periode: {data['period_label']}", sub_style))
    elements.append(Spacer(1, 6))

    if include_summary:
        elements.append(Paragraph("RINGKASAN PERIODE", sec_style))
        gt = data["grand_total"]
        sum_header = [Paragraph("Total Pegawai", table_header_style)] + [Paragraph(STATUS_LABEL[s], table_header_style) for s in STATUSES] + [Paragraph("Total Kehadiran", table_header_style)]
        sum_row = [Paragraph(str(data["total_pegawai"]), table_cell_center_bold)] + [Paragraph(str(gt[s]), table_cell_center_bold) for s in STATUSES] + [Paragraph(str(gt["HDR"]), table_cell_center_bold)]
        st = Table([sum_header, sum_row], repeatRows=1)
        st.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(st)
        elements.append(Spacer(1, 8))

        elements.append(Paragraph("REKAP PER PEGAWAI", sec_style))
        header = ["No", "NIP", "Nama", "Regu", "Kategori", "Jabatan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
        table_data = [[Paragraph(str(value), table_header_style) for value in header]]
        for row in data["rows"]:
            cells = [
                Paragraph(str(row["no"]), table_cell_center_style),
                Paragraph(str(row["nip"]), table_cell_center_style),
                Paragraph(str(row["nama"]), table_cell_nama_style),
                Paragraph(format_regu(row["regu"]), table_cell_center_style),
                Paragraph(str(row["category"]), table_cell_center_style),
                Paragraph(str(row["jabatan"]), table_cell_center_style),
            ]
            cells += [
                Paragraph(str(row[s]), status_cell_styles[s])
                for s in STATUSES
            ]
            cells += [
                Paragraph(str(row["jumlah_hari_kerja"]), table_cell_center_bold),
                Paragraph(str(row["total_kehadiran"]), table_cell_center_bold),
            ]
            table_data.append(cells)

        # Bottom of the Table (==== TOTAL ====)
        total_row = [
            Paragraph("TOTAL", table_cell_center_bold), "", "", "", "", ""
        ] + [
            Paragraph(str(data["grand_total"][s]), table_cell_center_bold)
            for s in STATUSES
        ] + [
            Paragraph(str(sum(row.get("jumlah_hari_kerja", 0) for row in data["rows"])), table_cell_center_bold),
            Paragraph(str(data["grand_total"]["HDR"]), table_cell_center_bold),
        ]
        table_data.append(total_row)
        
        colw = [8 * mm, 28 * mm, 42 * mm, 20 * mm, 20 * mm, 38 * mm] + [11 * mm] * 6 + [20 * mm, 20 * mm]
        t = Table(table_data, colWidths=colw, repeatRows=1)
        ts = [
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -2), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
            ("SPAN", (0, -1), (5, -1)),
        ]
        for i, s in enumerate(STATUSES):
            ts.append(("TEXTCOLOR", (6 + i, 1), (6 + i, -2), status_colors[s]))
        t.setStyle(TableStyle(ts))
        elements.append(t)

    if include_breakdown and len(data["breakdown"]) > 0:
        elements.append(Paragraph("BREAKDOWN BULANAN (mengikuti histori regu per bulan)", sec_style))
        header = ["No", "Nama", "Regu", "Bulan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
        bd = [[Paragraph(str(value), table_header_style) for value in header]]
        for b in data["breakdown"]:
            cells = [
                Paragraph(str(b["no"]), table_cell_center_style),
                Paragraph(str(b["nama"]), table_cell_nama_style),
                Paragraph(format_regu(b["regu"]), table_cell_center_style),
                Paragraph(str(b["month_label"]), table_cell_center_style),
            ]
            cells += [
                Paragraph(str(b[s]), status_cell_styles[s])
                for s in STATUSES
            ]
            cells += [
                Paragraph(str(b["jumlah_hari_kerja"]), table_cell_center_bold),
                Paragraph(str(b["total_kehadiran"]), table_cell_center_bold),
            ]
            bd.append(cells)
        colw = [8 * mm, 56 * mm, 26 * mm, 30 * mm] + [13 * mm] * 6 + [20 * mm, 20 * mm]
        t = Table(bd, colWidths=colw, repeatRows=1)
        ts = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (1, 1), (1, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for i, s in enumerate(STATUSES):
            ts.append(("TEXTCOLOR", (4 + i, 1), (4 + i, -1), status_colors[s]))
        t.setStyle(TableStyle(ts))
        elements.append(t)

    if include_detail:
        elements.append(Paragraph("DETAIL ABSENSI", sec_style))
        db = await get_db()
        ey, em = int(end[:4]), int(end[5:7])
        sd = f"{start}-01"
        ed = f"{end}-{last_day_of_month(ey, em):02d}"
        q = db.table("attendance").select("*").gte("date", sd).lte("date", ed).order("date")
        records = await fetch_all(q)
        emps = {e["id"]: e for e in ((await db.table("employees").select("*").execute()).data or [])}
        teams = {t["id"]: t for t in ((await db.table("teams").select("*").execute()).data or [])}
        dh_header = ["Tanggal", "NIP", "Nama", "Regu", "Status"]
        dh = [[Paragraph(str(value), table_header_style) for value in dh_header]]
        for rec in records:
            if team_id and rec.get("team_id") != team_id:
                continue
            e = emps.get(rec["employee_id"], {})
            dh.append([
                Paragraph(str(rec["date"]), table_cell_center_style),
                Paragraph(str(e.get("nip") or "-"), table_cell_center_style),
                Paragraph(str(e.get("nama") or "-"), table_cell_nama_style),
                Paragraph(str(teams.get(rec.get("team_id"), {}).get("name") or "-"), table_cell_center_style),
                Paragraph(str(rec["status"]), status_cell_styles.get(rec["status"], table_cell_center_bold)),
            ])
        t = Table(dh, colWidths=[26 * mm, 36 * mm, 70 * mm, 30 * mm, 20 * mm], repeatRows=1)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (2, 1), (2, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]))
        elements.append(t)

    doc.build(elements)
    buf.seek(0)
    fname = f"Rekap_Absensi_DAMKAR_{start}_{end}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )


@router.get("/export/kasubid/excel")
async def export_kasubid_excel(
    start: str,
    end: str,
    user: dict = Depends(require_roles("admin", "operator", "viewer", "kasubid")),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

    data = await _compute_kasubid_recap(start, end)
    wb = Workbook()
    thin = Side(style="thin", color="94A3B8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hf = PatternFill("solid", fgColor="0F172A")
    hfont = Font(bold=True, size=10, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)
    data_font = Font(size=10)
    data_font_bold = Font(size=10, bold=True)

    def head(ws, cols, row=1):
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=row, column=i, value=c)
            cell.fill = hf
            cell.font = hfont
            cell.alignment = center
            cell.border = border

    ws = wb.active
    ws.title = "Ringkasan Kasubid"
    cols = ["No", "Nama", "Posisi"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
    head(ws, cols)
    r = 2
    for i, p in enumerate(data["positions"], 1):
        vals = [i, p["nama"], p["label"]] + [p[s] for s in STATUSES] + [p["jumlah_hari_kerja"], p["total_kehadiran"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=sanitize_excel_cell(v))
            c.border = border
            c.font = data_font
            if j == 2:  # Nama: Left Center
                c.alignment = left_center
            else:       # Other: Middle
                c.alignment = center
            if 4 <= j <= 9:
                c.fill = PatternFill("solid", fgColor=STATUS_HEX[STATUSES[j - 4]])
                c.font = Font(size=10, color="FFFFFF", bold=True)
            elif j in (10, 11):
                c.font = data_font_bold
        r += 1

    ws.cell(row=r, column=1, value="TOTAL").font = data_font_bold
    ws.cell(row=r, column=1).alignment = center
    ws.cell(row=r, column=1).border = border
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    for j in range(2, 4):
        ws.cell(row=r, column=j).border = border
    for j, s in enumerate(STATUSES, 4):
        c = ws.cell(row=r, column=j, value=sum(p[s] for p in data["positions"]))
        c.font = data_font_bold
        c.alignment = center
        c.border = border
    c_jhk = ws.cell(row=r, column=10, value=sum(p.get("jumlah_hari_kerja", 0) for p in data["positions"]))
    c_jhk.font = data_font_bold
    c_jhk.alignment = center
    c_jhk.border = border
    c_hdr = ws.cell(row=r, column=11, value=sum(p["total_kehadiran"] for p in data["positions"]))
    c_hdr.font = data_font_bold
    c_hdr.alignment = center
    c_hdr.border = border

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 20
    for col in ["D", "E", "F", "G", "H", "I"]:
        ws.column_dimensions[col].width = 12
    for col in ["J", "K"]:
        ws.column_dimensions[col].width = 18

    ws2 = wb.create_sheet("Detail Bulanan")
    bcols = ["Nama", "Posisi", "Bulan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
    head(ws2, bcols)
    r = 2
    for p in data["positions"]:
        for mo in p["monthly"]:
            vals = [p["nama"], p["label"], mo["month_label"]] + [mo[s] for s in STATUSES] + [mo["jumlah_hari_kerja"], mo["total_kehadiran"]]
            for j, v in enumerate(vals, 1):
                c = ws2.cell(row=r, column=j, value=sanitize_excel_cell(v))
                c.border = border
                c.font = data_font
                if j == 1:  # Nama: Left Center
                    c.alignment = left_center
                else:       # Other: Middle
                    c.alignment = center
                if 4 <= j <= 9:
                    c.fill = PatternFill("solid", fgColor=STATUS_HEX[STATUSES[j - 4]])
                    c.font = Font(size=10, color="FFFFFF", bold=True)
                elif j >= 10:
                    c.font = data_font_bold
            r += 1
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 20
    ws2.column_dimensions["C"].width = 18
    for col in ["D", "E", "F", "G", "H", "I"]:
        ws2.column_dimensions[col].width = 12
    for col in ["J", "K"]:
        ws2.column_dimensions[col].width = 18

    ws3 = wb.create_sheet("Detail Harian")
    head(ws3, ["Tanggal", "Nama", "Posisi", "Status", "Keterangan"])
    r = 2
    for d in data["detail"]:
        vals = [d["date"], d["nama"], d["label"], d["status"], STATUS_LABEL.get(d["status"])]
        for j, v in enumerate(vals, 1):
            c = ws3.cell(row=r, column=j, value=sanitize_excel_cell(v))
            c.border = border
            c.font = data_font
            if j == 2:  # Nama: Left Center
                c.alignment = left_center
            else:       # Other: Middle
                c.alignment = center
            if j == 4:
                c.fill = PatternFill("solid", fgColor=STATUS_HEX.get(d["status"], "FFFFFF"))
                c.font = Font(size=10, color="FFFFFF", bold=True)
        r += 1
    for col in "ABCDE":
        ws3.column_dimensions[col].width = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=Rekap_Kasubid_{start}_{end}.xlsx"}
    )


@router.get("/export/kasubid/pdf")
async def export_kasubid_pdf(
    start: str,
    end: str,
    user: dict = Depends(require_roles("admin", "operator", "viewer", "kasubid")),
):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import BaseDocTemplate, PageTemplate, Frame, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT

    data = await _compute_kasubid_recap(start, end)
    buf = io.BytesIO()
    page = landscape(A4)
    doc = BaseDocTemplate(buf, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=30 * mm, bottomMargin=16 * mm)
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="m")
    gen = datetime.now().strftime("%d-%m-%Y %H:%M")

    def hf(canvas, d):
        canvas.saveState()
        w, h = page
        canvas.setFont("Helvetica-Bold", 12)
        canvas.drawCentredString(w / 2, h - 13 * mm, KOP["l1"])
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawCentredString(w / 2, h - 18 * mm, KOP["l2"])
        canvas.setFont("Helvetica", 8.5)
        canvas.drawCentredString(w / 2, h - 22 * mm, KOP["l3"])
        canvas.line(14 * mm, h - 25 * mm, w - 14 * mm, h - 25 * mm)
        canvas.setFont("Helvetica-Oblique", 7.5)
        canvas.drawString(14 * mm, 9 * mm, "Dicetak dari Sistem Informasi Absensi DAMKAR")
        canvas.drawRightString(w - 14 * mm, 9 * mm, f"Halaman {d.page}  |  {gen}")
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate(id="m", frames=[frame], onPage=hf)])
    styles = getSampleStyleSheet()
    title = ParagraphStyle("t", parent=styles["Title"], fontSize=13, alignment=TA_CENTER)
    sub = ParagraphStyle("s", parent=styles["Normal"], fontSize=9, alignment=TA_CENTER)
    sec_style = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=10, textColor=colors.HexColor("#0F172A"), spaceBefore=10, spaceAfter=4)
    
    CELL_FONT_SIZE = 7
    CELL_LEADING = 8.5

    table_header_style = ParagraphStyle(
        "table_header",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        textColor=colors.white,
        wordWrap="LTR",
    )

    table_cell_center_style = ParagraphStyle(
        "table_cell_center",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        wordWrap="LTR",
    )

    table_cell_center_bold = ParagraphStyle(
        "table_cell_center_bold",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_CENTER,
        wordWrap="LTR",
    )

    table_cell_nama_style = ParagraphStyle(
        "table_cell_nama",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=CELL_FONT_SIZE,
        leading=CELL_LEADING,
        alignment=TA_LEFT,
        wordWrap="LTR",
    )

    status_colors = {s: colors.HexColor("#" + STATUS_HEX[s]) for s in STATUSES}

    status_cell_styles = {
        s: ParagraphStyle(
            f"table_status_{s}",
            parent=styles["Normal"],
            textColor=status_colors[s],
            fontName="Helvetica-Bold",
            fontSize=CELL_FONT_SIZE,
            leading=CELL_LEADING,
            alignment=TA_CENTER,
            wordWrap="LTR",
        )
        for s in STATUSES
    }

    el = [
        Paragraph("REKAP ABSENSI KASUBID", title),
        Paragraph(f"Periode: {data['period_label']}", sub),
        Spacer(1, 8)
    ]
    header = ["No", "Nama", "Posisi"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
    td = [[Paragraph(str(value), table_header_style) for value in header]]
    for i, p in enumerate(data["positions"], 1):
        cells = [
            Paragraph(str(i), table_cell_center_style),
            Paragraph(p["nama"], table_cell_nama_style),
            Paragraph(p["label"], table_cell_center_style),
        ]
        cells += [
            Paragraph(str(p[s]), status_cell_styles[s])
            for s in STATUSES
        ]
        cells += [
            Paragraph(str(p["jumlah_hari_kerja"]), table_cell_center_bold),
            Paragraph(str(p["total_kehadiran"]), table_cell_center_bold),
        ]
        td.append(cells)

    total_row = [
        Paragraph("TOTAL", table_cell_center_bold), "", ""
    ] + [
        Paragraph(str(sum(p[s] for p in data["positions"])), table_cell_center_bold)
        for s in STATUSES
    ] + [
        Paragraph(str(sum(p.get("jumlah_hari_kerja", 0) for p in data["positions"])), table_cell_center_bold),
        Paragraph(str(sum(p["total_kehadiran"] for p in data["positions"])), table_cell_center_bold),
    ]
    td.append(total_row)

    colw = [12 * mm, 75 * mm, 32 * mm] + [14 * mm] * 6 + [20 * mm, 20 * mm]
    t = Table(td, repeatRows=1, colWidths=colw)
    ts = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("ALIGN", (1, 1), (1, -2), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E2E8F0")),
        ("SPAN", (0, -1), (2, -1)),
    ]
    for idx, s in enumerate(STATUSES):
        ts.append(("TEXTCOLOR", (3 + idx, 1), (3 + idx, -2), status_colors[s]))
    t.setStyle(TableStyle(ts))
    el.append(t)

    if len(data["months"]) > 1:
        el.append(Spacer(1, 10))
        el.append(Paragraph("DETAIL BULANAN", sec_style))
        b_header = ["Nama", "Posisi", "Bulan"] + STATUSES + ["Jumlah Hari Kerja", "Total Kehadiran"]
        b_td = [[Paragraph(str(value), table_header_style) for value in b_header]]
        for p in data["positions"]:
            for mo in p["monthly"]:
                b_cells = [
                    Paragraph(p["nama"], table_cell_nama_style),
                    Paragraph(p["label"], table_cell_center_style),
                    Paragraph(mo["month_label"], table_cell_center_style),
                ]
                b_cells += [Paragraph(str(mo[s]), status_cell_styles[s]) for s in STATUSES]
                b_cells += [
                    Paragraph(str(mo["jumlah_hari_kerja"]), table_cell_center_bold),
                    Paragraph(str(mo["total_kehadiran"]), table_cell_center_bold),
                ]
                b_td.append(b_cells)
        b_colw = [70 * mm, 32 * mm, 28 * mm] + [14 * mm] * 6 + [20 * mm, 20 * mm]
        b_t = Table(b_td, repeatRows=1, colWidths=b_colw)
        b_ts = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("ALIGN", (0, 1), (0, -1), "LEFT"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ]
        for idx, s in enumerate(STATUSES):
            b_ts.append(("TEXTCOLOR", (3 + idx, 1), (3 + idx, -1), status_colors[s]))
        b_t.setStyle(TableStyle(b_ts))
        el.append(b_t)

    doc.build(el)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=Rekap_Kasubid_{start}_{end}.pdf"}
    )
